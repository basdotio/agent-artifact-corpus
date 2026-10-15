#!/usr/bin/env python3
"""
Extract plain text from an existing resume (PDF, DOCX, TXT, MD) so it can seed
the career profile.

Handles what designed resumes hide from simple readers: DOCX tables, text in
headers/footers, and text boxes; PDFs in reading order (-layout off) and with
layout (-layout on) so two-column designs can be untangled. Also prints a short
parse report: likely contact info, dates found, and warnings (image-only PDF,
ligature glyphs, columns, tables).

Usage:
    extract.py old-resume.pdf [--out profile/sources/old-resume.txt]
    extract.py old-resume.docx
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def ptext(p) -> str:
    """Paragraph text keeping tabs and line breaks (w:tab, w:br, w:cr)."""
    out = []
    for el in p.iter():
        if el.tag == W + "t":
            out.append(el.text or "")
        elif el.tag == W + "tab":
            out.append("\t")
        elif el.tag in (W + "br", W + "cr"):
            out.append("\n")
    return "".join(out)


def from_docx(path: Path, warns: list) -> str:
    from docx import Document

    d = Document(str(path))
    parts = []
    hf = []
    for s in d.sections:
        for part in (s.header, s.first_page_header):
            t = "\n".join(p.text for p in part.paragraphs).strip()
            if t and t not in hf:
                hf.append(t)
    if hf:
        warns.append("contact/info in Word header (many ATS drop it); recovered below")
        parts.append("\n".join(hf))
    # walk the body in order, including tables and text boxes
    body = d.element.body
    n_tables = 0
    for el in body.iterchildren():
        if el.tag == W + "p":
            parts.append(ptext(el))
        elif el.tag == W + "tbl":
            n_tables += 1
            for row in el.iter(W + "tr"):
                cells = []
                for cell in row.iter(W + "tc"):
                    ps = [ptext(p) for p in cell.iter(W + "p")]
                    cells.append("\n".join(x for x in ps if x.strip()))
                parts.append("\n".join(c for c in cells if c.strip()))
    if n_tables:
        warns.append(f"{n_tables} table(s) in layout (ATS often scrambles these); flattened cell by cell")
    boxes = len(list(body.iter(W + "txbxContent")))
    if boxes:
        warns.append(f"{boxes} text box(es) found (ATS usually skips these)")
    for s in d.sections:
        t = "\n".join(p.text for p in s.footer.paragraphs).strip()
        if t:
            parts.append(t)
    return "\n".join(parts)


def from_pdf(path: Path, warns: list) -> str:
    if shutil.which("pdftotext"):
        raw = subprocess.run(["pdftotext", "-enc", "UTF-8", str(path), "-"], capture_output=True, text=True).stdout
        lay = subprocess.run(["pdftotext", "-enc", "UTF-8", "-layout", str(path), "-"],
                             capture_output=True, text=True).stdout
    else:
        from pypdf import PdfReader

        raw = lay = "\n".join(p.extract_text() or "" for p in PdfReader(str(path)).pages)
    if len(raw.split()) < 40:
        warns.append("little or no text layer: likely an image/scanned PDF. Read the PDF visually instead, "
                     "or ask the user for the DOCX/original")
    # two-column hint: many lines with a wide internal gap in layout mode
    gappy = sum(1 for l in lay.splitlines() if re.search(r"\S {6,}\S", l) and not re.search(
        r"\b(19|20)\d{2}\s*$|Present\s*$", l))
    if gappy > 8:
        warns.append(f"looks multi-column ({gappy} split lines); compare both texts below, reading order may be mixed")
        return raw + "\n\n===== LAYOUT VIEW (for untangling columns) =====\n" + lay
    return raw


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--out", type=Path)
    a = ap.parse_args()

    warns: list[str] = []
    ext = a.file.suffix.lower()
    if ext == ".pdf":
        text = from_pdf(a.file, warns)
    elif ext == ".docx":
        text = from_docx(a.file, warns)
    elif ext in (".txt", ".md"):
        text = a.file.read_text(encoding="utf-8", errors="replace")
    elif ext == ".doc":
        sys.exit("Old .doc format: ask the user to re-save as .docx or PDF (or `textutil -convert txt` on macOS).")
    else:
        sys.exit(f"Unsupported file type: {ext}")

    if re.search("[ﬀ-ﬆ]", text):
        warns.append("ligature glyphs found (e.g. 'ﬁ'); keyword matching on the old file was likely broken")
        text = (text.replace("ﬁ", "fi").replace("ﬂ", "fl").replace("ﬀ", "ff")
                .replace("ﬃ", "ffi").replace("ﬄ", "ffl"))
    zw = sum(text.count(c) for c in "\u200b\u200c\u200d\ufeff")
    if zw:
        warns.append(f"{zw} zero-width characters (Google Docs adds these); they can split words for keyword search. Removed.")
        text = re.sub("[\u200b\u200c\u200d\ufeff]", "", text)
    if re.search(r"\w[\u2013\u2014]", text):
        warns.append("dash glued to a word (e.g. 'Engineer\u2013 December'); parsers may read the title and date as one token")
    text = re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"

    if a.out:
        a.out.parent.mkdir(parents=True, exist_ok=True)
        a.out.write_text(text, encoding="utf-8")

    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text)
    phone = re.search(r"\(?\d{3}\)?[ .-]?\d{3}[ .-]?\d{4}", text)
    linkedin = re.search(r"linkedin\.com/in/[\w-]+", text, re.I)
    dates = re.findall(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.? \d{4}|\b\d{1,2}/\d{4}\b", text)
    print("=== PARSE REPORT ===")
    print(f"words: {len(text.split())}")
    print(f"email: {email.group(0) if email else 'NOT FOUND'}")
    print(f"phone: {phone.group(0) if phone else 'NOT FOUND'}")
    print(f"linkedin: {linkedin.group(0) if linkedin else 'not found'}")
    print(f"dates found: {len(dates)}")
    for w in warns:
        print(f"WARN: {w}")
    if a.out:
        print(f"saved: {a.out}")
    print("=== TEXT ===")
    print(text)


if __name__ == "__main__":
    main()
