#!/usr/bin/env python3
"""
Verify that generated resume/cover-letter files are machine readable and that
JD keywords survive extraction.

Checks (each prints PASS/WARN/FAIL):
  - text layer extracts from PDF (pdftotext) and DOCX (python-docx)
  - name on the first line, email and phone found near the top
  - nothing in DOCX headers/footers, no hidden/white/tiny text
  - standard section headings present, Summary first, Skills last (resume)
  - no ligature glyphs, smart quotes, em/en dashes, zero-width or odd unicode
  - consistent date format on role lines
  - fonts embedded, file size < 2.5 MB, page count within limit
  - resume word count within budget (default 475-750; --min-words / --max-words)
  - keyword coverage by section (with --keywords)

Keyword file: one term per line, "must:" or "nice:" prefix, alternates after "|".
    must: Kubernetes | k8s
    must: Amazon Web Services | AWS
    nice: Terraform

Usage:
    verify.py --source resume.md --pdf R.pdf --docx R.docx [--keywords kw.txt] [--max-pages 2]
    verify.py --type cover --source cover-letter.md --pdf C.pdf --docx C.docx
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

STANDARD = {
    "summary", "professional summary", "experience", "work experience", "professional experience",
    "projects", "education", "skills", "technical skills", "certifications", "publications",
    "patents", "awards", "volunteer experience", "open source", "speaking",
}
BAD_CHARS = {
    "—": "em dash", "–": "en dash", "‘": "curly quote", "’": "curly quote",
    "“": "curly quote", "”": "curly quote", "…": "ellipsis char",
    "​": "zero-width space", "‌": "zero-width non-joiner", "‍": "zero-width joiner",
    "﻿": "BOM", " ": "non-breaking space", "­": "soft hyphen",
}
LIGATURES = re.compile("[ﬀ-ﬆ]")
MONTHS = r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
DATE_RANGE = re.compile(rf"^{MONTHS} \d{{4}} - ({MONTHS} \d{{4}}|Present)$")
YEAR_ONLY = re.compile(r"^(\d{4}|\d{4} - (\d{4}|Present))$")

results: list[tuple[str, str]] = []


def report(level: str, msg: str):
    results.append((level, msg))


def pdf_text(pdf: Path) -> str:
    if not shutil.which("pdftotext"):
        from pypdf import PdfReader

        return "\n".join(p.extract_text() or "" for p in PdfReader(str(pdf)).pages)
    return subprocess.run(["pdftotext", "-enc", "UTF-8", str(pdf), "-"], capture_output=True, text=True).stdout


def check_pdf(pdf: Path, max_pages: int) -> str:
    from pypdf import PdfReader

    size = pdf.stat().st_size
    report("PASS" if size < 2_500_000 else "FAIL", f"PDF size {size/1024:.0f} KB (limit 2.5 MB)")
    r = PdfReader(str(pdf))
    n = len(r.pages)
    report("PASS" if n <= max_pages else "FAIL", f"PDF pages: {n} (max {max_pages})")
    if r.is_encrypted:
        report("FAIL", "PDF is encrypted/password protected")
    unembedded = set()
    for page in r.pages:
        fonts = (page.get("/Resources") or {}).get("/Font") or {}
        for f in fonts.values():
            f = f.get_object()
            desc = f.get("/FontDescriptor")
            if f.get("/Subtype") == "/Type0":
                desc = f["/DescendantFonts"][0].get_object().get("/FontDescriptor")
            if desc is None or not any(k in desc.get_object() for k in ("/FontFile", "/FontFile2", "/FontFile3")):
                unembedded.add(str(f.get("/BaseFont")))
    report("PASS" if not unembedded else "WARN",
           "PDF fonts embedded" if not unembedded else f"PDF fonts not embedded: {sorted(unembedded)}")
    text = pdf_text(pdf)
    report("PASS" if len(text.split()) > 30 else "FAIL", f"PDF text layer: {len(text.split())} words extracted")
    return text


def check_docx(path: Path) -> str:
    from docx import Document

    d = Document(str(path))
    size = path.stat().st_size
    report("PASS" if size < 2_500_000 else "FAIL", f"DOCX size {size/1024:.0f} KB")
    hf = []
    for s in d.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            t = "\n".join(p.text for p in part.paragraphs).strip()
            if t:
                hf.append(t)
    report("PASS" if not hf else "FAIL",
           "DOCX header/footer empty" if not hf else f"DOCX header/footer has text (ATS often drops it): {hf}")
    if d.tables:
        report("FAIL", f"DOCX contains {len(d.tables)} table(s)")
    hidden = []
    for p in d.paragraphs:
        for r in p.runs:
            rpr = r._r.rPr
            vanish = rpr is not None and rpr.find(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}vanish") is not None
            white = r.font.color is not None and r.font.color.rgb is not None and str(r.font.color.rgb) in ("FFFFFF", "FEFEFE")
            tiny = r.font.size is not None and r.font.size.pt < 7
            if (vanish or white or tiny) and r.text.strip():
                hidden.append(r.text.strip()[:40])
    report("PASS" if not hidden else "FAIL",
           "DOCX has no hidden/white/tiny text" if not hidden else f"DOCX hidden text found: {hidden}")
    text = "\n".join(p.text for p in d.paragraphs)
    report("PASS" if len(text.split()) > 30 else "FAIL", f"DOCX text: {len(text.split())} words")
    return text


def check_text(label: str, text: str, kind: str):
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    head = "\n".join(lines[:4])
    if kind == "resume" or kind == "cover":
        report("PASS" if re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", head) else "FAIL", f"{label}: email in top lines")
        report("PASS" if re.search(r"\(?\d{3}\)?[ .-]?\d{3}[ .-]?\d{4}", head) else "WARN", f"{label}: phone in top lines")
    lig = LIGATURES.findall(text)
    report("PASS" if not lig else "FAIL", f"{label}: no ligature glyphs" if not lig else f"{label}: ligatures {set(lig)}")
    bad = {}
    for ch, name in BAD_CHARS.items():
        if ch in text:
            bad[name] = bad.get(name, 0) + text.count(ch)
    odd = {c for c in text if ord(c) > 127 and c not in BAD_CHARS and c not in "•éèáñüöäç"}
    report("PASS" if not bad else "WARN", f"{label}: clean punctuation" if not bad else f"{label}: {bad}")
    if odd:
        report("WARN", f"{label}: unusual characters {sorted(odd)}")
    return lines


def check_source(src: Path, kind: str, keywords: Path | None, max_pages: int):
    text = re.sub(r"<!--.*?-->", "", src.read_text(encoding="utf-8"), flags=re.S)
    if kind == "cover":
        return text
    heads = [h.strip() for h in re.findall(r"^## (.+)$", text, re.M)]
    nonstd = [h for h in heads if h.lower() not in STANDARD]
    report("PASS" if not nonstd else "WARN",
           "section headings standard" if not nonstd else f"non-standard headings (parsers may miss): {nonstd}")
    if heads:
        report("PASS" if heads[0].lower() in ("summary", "professional summary") else "WARN",
               f"first section is '{heads[0]}' (expected Summary)")
        report("PASS" if "skill" in heads[-1].lower() else "WARN", f"last section is '{heads[-1]}' (expected Skills)")
    m = re.search(r"^## (?:Professional )?Summary\s*\n(.+?)(?=^## )", text, re.M | re.S)
    if m:
        summ = " ".join(m.group(1).split())
        n = len(re.findall(r"[.!?](?:\s|$)", summ))
        report("PASS" if 3 <= n <= 4 else "WARN", f"summary has {n} sentences (target 3-4)")
    roles = re.findall(r"^### (.+)$", text, re.M)
    bad_dates = []
    for r in roles:
        parts = [p.strip() for p in r.split("|")]
        d = parts[3] if len(parts) > 3 else ""
        if d and not (DATE_RANGE.match(d) or YEAR_ONLY.match(d)):
            bad_dates.append(d)
        if len(parts) < 4:
            bad_dates.append(f"(missing field) {r}")
    report("PASS" if not bad_dates else "WARN",
           "role dates use 'Mon YYYY - Mon YYYY|Present'" if not bad_dates else f"date format issues: {bad_dates}")
    return text


def section_map(md: str) -> dict[str, str]:
    out, cur = {}, "header"
    for line in md.splitlines():
        if line.startswith("## "):
            cur = line[3:].strip().lower()
            continue
        out[cur] = out.get(cur, "") + line + "\n"
    return out


def term_re(t: str) -> re.Pattern:
    t = t.strip()
    pre = r"(?<![A-Za-z0-9])" if t[0].isalnum() else ""
    post = r"(?:s|es)?(?![A-Za-z0-9])" if t[-1].isalnum() else ""  # allow plurals
    return re.compile(pre + re.escape(t) + post, re.I)


def check_keywords(kw: Path, md: str, extracted: list[tuple[str, str]]):
    secs = section_map(md)
    rows, must_total, must_hit, nice_total, nice_hit = [], 0, 0, 0, 0
    skill_only = []
    for raw in kw.read_text(encoding="utf-8").splitlines():
        raw = raw.strip()
        if not raw or raw.startswith("#"):
            continue
        prio = "must"
        m = re.match(r"^(must|nice)\s*:\s*(.+)$", raw, re.I)
        if m:
            prio, raw = m.group(1).lower(), m.group(2)
        terms = [t.strip() for t in raw.split("|") if t.strip()]
        where = []
        for name, body in secs.items():
            if any(term_re(t).search(body) for t in terms):
                short = "summary" if "summary" in name else "skills" if "skill" in name else \
                        "experience" if "experience" in name else name
                where.append(short)
        exact = any(term_re(terms[0]).search(md) for _ in [0])
        lost = [lab for lab, txt in extracted if where and not any(term_re(t).search(txt) for t in terms)]
        hit = bool(where)
        if prio == "must":
            must_total += 1
            must_hit += hit
        else:
            nice_total += 1
            nice_hit += hit
        if where == ["skills"]:
            skill_only.append(terms[0])
        flag = "ok" if hit else "MISSING"
        if hit and not exact:
            flag = "alt-only"
        if lost:
            flag += f" LOST-IN-{','.join(lost)}"
        rows.append((prio, terms[0], flag, ",".join(sorted(set(where))) or "-"))
    print("\nKeyword coverage (term | priority | status | sections)")
    for prio, t, flag, where in rows:
        print(f"  {t:<38} {prio:<5} {flag:<14} {where}")
    if must_total:
        pct = 100 * must_hit / must_total
        report("PASS" if pct >= 85 else "WARN", f"must-have coverage {must_hit}/{must_total} ({pct:.0f}%)")
    if nice_total:
        report("PASS", f"nice-to-have coverage {nice_hit}/{nice_total}")
    if skill_only:
        report("WARN", "only in Skills, no supporting bullet (AI screeners may mark 'unknown'): " + ", ".join(skill_only))
    if any("alt-only" in r[2] for r in rows):
        report("WARN", "some terms appear only as alternates; use the JD's exact spelling at least once: "
               + ", ".join(r[1] for r in rows if "alt-only" in r[2]))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--type", choices=["resume", "cover"], default="resume")
    ap.add_argument("--source", type=Path, required=True)
    ap.add_argument("--pdf", type=Path)
    ap.add_argument("--docx", type=Path)
    ap.add_argument("--keywords", type=Path)
    ap.add_argument("--max-pages", type=int, default=None)
    ap.add_argument("--min-words", type=int, default=475)
    ap.add_argument("--max-words", type=int, default=750)
    a = ap.parse_args()
    max_pages = a.max_pages or (1 if a.type == "cover" else 2)

    md = check_source(a.source, a.type, a.keywords, max_pages)
    extracted = []
    if a.pdf:
        t = check_pdf(a.pdf, max_pages)
        check_text("PDF", t, a.type)
        extracted.append(("PDF", t))
        first = next((l for l in t.splitlines() if l.strip()), "")
        name = re.search(r"^# (.+)$", md, re.M)
        if name:
            report("PASS" if name.group(1).strip().lower() in first.lower() else "FAIL",
                   f"PDF first line is the name ('{first.strip()[:40]}')")
    if a.docx:
        t = check_docx(a.docx)
        check_text("DOCX", t, a.type)
        extracted.append(("DOCX", t))
    if a.type == "resume" and extracted:
        words = len(extracted[0][1].split())
        ok = a.min_words <= words <= a.max_words
        report("PASS" if ok else "WARN",
               f"resume length {words} words (budget {a.min_words}-{a.max_words})" if ok else
               f"resume length {words} words, outside {a.min_words}-{a.max_words}: "
               + ("cut the weakest and oldest bullets" if words > a.max_words else "add the strongest missing results"))
    if a.keywords and a.type == "resume":
        check_keywords(a.keywords, md, extracted)

    print()
    for level, msg in results:
        print(f"[{level}] {msg}")
    fails = sum(1 for l, _ in results if l == "FAIL")
    warns = sum(1 for l, _ in results if l == "WARN")
    print(f"\n{fails} FAIL, {warns} WARN")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
