"""Preview-first local file imports with explicit mapping and atomic unique-key upserts."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import uuid
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from .schema import META
from .store import Hub, now


def read_file(path: str | Path, sheet: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    path = Path(path)
    extension = path.suffix.lower()
    if extension == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            headers = reader.fieldnames or []
            rows = list(reader)
    elif extension == ".json":
        rows = json.loads(path.read_text(encoding="utf-8-sig"))
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError("JSON import must be an array of objects")
        headers = list(dict.fromkeys(key for row in rows for key in row))
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is not None and sheet not in workbook.sheetnames:
                raise ValueError("requested worksheet was not found")
            worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
            values = worksheet.iter_rows(values_only=True)
            headers = [
                str(value).strip() if value is not None else "" for value in next(values, [])
            ]
            rows = [
                dict(zip(headers, row, strict=False))
                for row in values
                if any(value is not None for value in row)
            ]
        finally:
            workbook.close()
    elif extension == ".pdf":
        import pdfplumber

        headers, rows = [], []
        with pdfplumber.open(path) as document:
            for page in document.pages:
                for table in page.extract_tables():
                    if not table:
                        continue
                    current = [str(cell or "").strip() for cell in table[0]]
                    if not headers:
                        headers = current
                    elif current != headers:
                        raise ValueError(
                            "PDF tables have inconsistent headers; import a normalized CSV instead"
                        )
                    for row in table[1:]:
                        if any(cell for cell in row) and row != headers:
                            rows.append(dict(zip(headers, row, strict=False)))
        if not headers:
            raise ValueError(
                "PDF has no extractable ruled tables; scanned PDFs require external OCR first"
            )
    else:
        raise ValueError("supported file types: .xlsx .csv .json .pdf")
    if (
        not headers
        or any(not isinstance(h, str) or not h for h in headers)
        or len(headers) != len(set(headers))
    ):
        raise ValueError("input needs unique non-empty column headers")
    if any(None in row for row in rows):
        raise ValueError("CSV rows contain more values than headers")
    return headers, rows


def normalized(value: str) -> str:
    return re.sub(r"[^\w]", "", value.casefold())


def suggest_mapping(headers: list[str], schema: dict[str, Any]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used = set()
    for header in headers:
        if header in META:
            continue
        candidates = []
        for name, options in schema["columns"].items():
            aliases = [name, options.get("label", name), *options.get("aliases", [])]
            score = max(
                SequenceMatcher(None, normalized(header), normalized(alias)).ratio()
                for alias in aliases
            )
            candidates.append((score, name))
        score, name = max(candidates)
        if score >= 0.72 and name not in used:
            mapping[header] = name
            used.add(name)
    return mapping


def preview(
    hub: Hub,
    path: str | Path,
    table: str,
    mapping: dict[str, str] | None = None,
    sheet: str | None = None,
    reveal: bool = False,
) -> dict[str, Any]:
    headers, raw = read_file(path, sheet)
    schema = hub._schema(table)
    mapping = suggest_mapping(headers, schema) if mapping is None else mapping
    if not isinstance(mapping, dict) or any(
        source not in headers or target not in schema["columns"]
        for source, target in mapping.items()
    ):
        raise ValueError("mapping must map existing source headers to declared business fields")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("mapping cannot assign multiple source columns to the same field")
    rows = [{target: row.get(source) for source, target in mapping.items()} for row in raw]
    errors = []
    validated = []
    for index, row in enumerate(rows, start=2):
        try:
            validated.append(hub._prepare(table, row))
        except ValueError as exc:
            errors.append({"row": index, "message": str(exc)})
    unmapped = [column for column in headers if column not in mapping and column not in META]
    return {
        "table": table,
        "file": Path(path).name,
        "headers": headers,
        "mapping": mapping,
        "unmapped_columns": unmapped,
        "rows_in": len(rows),
        "valid_rows": len(validated),
        "errors": errors,
        "sample": [hub.mask(table, row, reveal) for row in validated[:5]],
        "requires_confirmation": True,
        "sha1": hashlib.sha1(Path(path).read_bytes()).hexdigest(),
    }


def apply_import(
    hub: Hub,
    path: str | Path,
    table: str,
    mapping: dict[str, str] | None = None,
    sheet: str | None = None,
    confirm: bool = False,
    actor: str = "cli",
    expected_sha1: str | None = None,
) -> dict[str, Any]:
    """Write one validated import transaction; repeated identical inputs produce no row changes."""
    if confirm is not True:
        raise ValueError("import writes require confirm=true or --yes; run a preview first")
    plan = preview(hub, path, table, mapping, sheet)
    if expected_sha1 is not None and plan["sha1"] != expected_sha1:
        raise ValueError("file changed since preview; preview it again")
    if plan["errors"]:
        raise ValueError(
            f"import validation failed in {len(plan['errors'])} row(s); run --dry-run for details"
        )
    schema = hub._schema(table)
    if not schema["unique"]:
        raise ValueError("import requires schema unique keys for deterministic upserts")
    _, rows = read_file(path, sheet)
    inserted = updated = skipped = 0
    with hub.connect(write=True) as conn:
        for raw in rows:
            data = {target: raw.get(source) for source, target in plan["mapping"].items()}
            prepared = hub._prepare(table, data)
            keys = schema["unique"]
            clause = " AND ".join('"' + name + '"=?' for name in keys)
            existing = conn.execute(
                f'SELECT * FROM "h_{table}" WHERE {clause}', [prepared[name] for name in keys]
            ).fetchone()
            if existing:
                # Only mapped columns are replaced; fields omitted by a source survive.
                _, change = hub._update(conn, table, existing["id"], data, actor)
                if change:
                    updated += 1
                else:
                    skipped += 1
            else:
                row_id = raw.get("id")
                if row_id is not None and (not isinstance(row_id, str) or not row_id):
                    raise ValueError("imported id must be a non-empty string")
                hub._insert(conn, table, data, actor, Path(path).name, row_id=row_id)
                inserted += 1
        import_id = str(uuid.uuid4())
        conn.execute(
            "INSERT INTO imports VALUES(?,?,?,?,?,?,?)",
            (import_id, Path(path).name, plan["sha1"], table, len(rows), inserted + updated, now()),
        )
    return {
        "import_id": import_id,
        "rows_in": len(rows),
        "inserted": inserted,
        "updated": updated,
        "skipped": skipped,
        "sha1": plan["sha1"],
    }
