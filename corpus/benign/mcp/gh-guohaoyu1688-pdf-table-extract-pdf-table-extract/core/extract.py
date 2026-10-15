"""Table extraction with explicit geometry, provenance and review-oriented quality reports."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections.abc import Callable
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import pdfplumber
import pypdfium2 as pdfium
from jsonschema import Draft202012Validator
from openpyxl import Workbook

from .errors import PDFError

RESULT_VALIDATOR = Draft202012Validator(
    json.loads(Path(__file__).with_name("result.schema.json").read_text(encoding="utf-8"))
)


def selected_pages(value: str | None, count: int) -> list[int]:
    if value is None or value == "all":
        return list(range(1, count + 1))
    result = set()
    try:
        for part in value.split(","):
            pieces = part.strip().split("-")
            if len(pieces) == 1:
                result.add(int(pieces[0]))
            elif len(pieces) == 2:
                start, end = map(int, pieces)
                if start > end:
                    raise ValueError()
                if start < 1 or end > count:
                    raise ValueError()
                result.update(range(start, end + 1))
            else:
                raise ValueError()
        if not result or min(result) < 1 or max(result) > count:
            raise ValueError()
    except (ValueError, TypeError) as exc:
        raise PDFError("invalid_pages", f"Use page numbers/ranges within 1..{count}", 2) from exc
    return sorted(result)


def open_pdf(path: str | Path):
    path = Path(path).expanduser()
    if not path.is_file():
        raise PDFError("not_found", f"PDF not found: {path}", 3)
    try:
        return pdfplumber.open(path, unicode_norm="NFC")
    except Exception as exc:
        raise PDFError(
            "unreadable_pdf", "Cannot open PDF; encrypted files need an unlocked copy"
        ) from exc


def page_area(page: Any, area: list[float] | None, trim_top: float, trim_bottom: float):
    bounds = area or [0, trim_top, float(page.width), float(page.height) - trim_bottom]
    if (
        len(bounds) != 4
        or any(not isinstance(v, (int, float)) for v in bounds)
        or not 0 <= bounds[0] < bounds[2] <= page.width
        or not 0 <= bounds[1] < bounds[3] <= page.height
    ):
        raise PDFError("invalid_area", "Area must be x0,top,x1,bottom inside the page", 2)
    return page.crop(tuple(bounds)) if tuple(bounds) != page.bbox else page


def plumber_tables(page: Any, columns: list[float] | None = None) -> list[dict]:
    """Prefer ruled tables, then fall back to aligned-text table detection."""
    if columns:
        if len(columns) < 2 or sorted(set(columns)) != columns:
            raise PDFError("invalid_columns", "Columns need two or more increasing x boundaries", 2)
        if columns[0] < page.bbox[0] or columns[-1] > page.bbox[2]:
            raise PDFError(
                "invalid_columns", "Column boundaries must lie within the selected area", 2
            )
    methods = ["lines", "text"]
    for strategy in methods:
        settings = {
            "vertical_strategy": "explicit" if columns else strategy,
            "horizontal_strategy": strategy,
            "snap_tolerance": 3,
            "intersection_tolerance": 5,
            "min_words_vertical": 2,
            "min_words_horizontal": 1,
        }
        if columns:
            settings["explicit_vertical_lines"] = columns
        found = page.find_tables(settings)
        tables = []
        for table in found:
            raw = table.extract(x_tolerance=3, y_tolerance=3)
            cells = [list(cell) for cell in table.cells if cell]
            notes = []
            raw = [row for row in raw if any(value not in (None, "") for value in row)]
            if len(raw) < 2 or max(map(len, raw), default=0) < 2:
                continue
            # Text alignment can include isolated page headings/footers outside the table.
            if strategy == "text":
                raw = [row for row in raw if sum(bool(value) for value in row) >= 2]
                if len(raw) < 2:
                    continue
                if not columns and any(not value for value in raw[0]) and raw[0][0]:
                    # Word alignment may split "Sample A" into two columns, while
                    # the header correctly names the logical column only once.
                    groups = []
                    for index, header in enumerate(raw[0]):
                        if header or not groups:
                            groups.append([index])
                        else:
                            groups[-1].append(index)
                    raw = [
                        [" ".join(row[i] or "" for i in group).strip() for group in groups]
                        for row in raw
                    ]
                    cells = []
                    for row in table.rows:
                        for group in groups:
                            pieces = [row.cells[i] for i in group if row.cells[i]]
                            if pieces:
                                cells.append(
                                    [
                                        min(c[0] for c in pieces),
                                        min(c[1] for c in pieces),
                                        max(c[2] for c in pieces),
                                        max(c[3] for c in pieces),
                                    ]
                                )
                    notes.append(
                        "Merged adjacent text fragments under an unnamed header; "
                        "verify column grouping."
                    )
            tables.append(
                {
                    "raw_rows": raw,
                    "bbox": list(table.bbox),
                    "cells": cells,
                    "engine": f"plumber-{strategy}",
                    "notes": notes,
                }
            )
        if tables:
            return tables
    return []


def render_page(path: str | Path, page: int = 1, scale: float = 1.5) -> bytes:
    if not 0.25 <= scale <= 4:
        raise PDFError("invalid_scale", "scale must be between 0.25 and 4", 2)
    if not Path(path).is_file():
        raise PDFError("not_found", f"PDF not found: {path}", 3)
    try:
        opened = pdfium.PdfDocument(str(path))
    except Exception as exc:
        raise PDFError(
            "unreadable_pdf", "Cannot render PDF; provide a readable unlocked copy"
        ) from exc
    with opened as document:
        if not 1 <= page <= len(document):
            raise PDFError("invalid_pages", "Page does not exist", 2)
        rendered_page = document[page - 1]
        try:
            bitmap = rendered_page.render(scale=scale)
            try:
                image = bitmap.to_pil()
                output = io.BytesIO()
                image.save(output, format="PNG")
                return output.getvalue()
            finally:
                bitmap.close()
        finally:
            rendered_page.close()


def ocr_tables(
    path: Path, page_number: int, area: list[float] | None, columns: list[float] | None
) -> list[dict]:
    """Reconstruct simple aligned tables from RapidOCR word boxes; models are bundled."""
    try:
        import numpy as np
        from PIL import Image
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise PDFError(
            "missing_dependency", "Install OCR support: pip install 'pdf-table-extract[ocr]'", 3
        ) from exc
    scale = 2.5
    image = Image.open(io.BytesIO(render_page(path, page_number, scale)))
    engine = RapidOCR(intra_op_num_threads=1, inter_op_num_threads=1)
    result, _ = engine(np.array(image))
    boxes = []
    for polygon, text, score in result or []:
        xs, ys = zip(*polygon)
        box = [min(xs) / scale, min(ys) / scale, max(xs) / scale, max(ys) / scale]
        if area and not (
            area[0] <= (box[0] + box[2]) / 2 <= area[2]
            and area[1] <= (box[1] + box[3]) / 2 <= area[3]
        ):
            continue
        boxes.append({"box": box, "text": text, "score": float(score)})
    if not boxes:
        return []
    rows: list[list[dict]] = []
    for box in sorted(boxes, key=lambda b: (b["box"][1] + b["box"][3]) / 2):
        center = (box["box"][1] + box["box"][3]) / 2
        if rows:
            last = rows[-1][0]["box"]
            if abs(center - (last[1] + last[3]) / 2) <= max(4, (last[3] - last[1]) * 0.7):
                rows[-1].append(box)
                continue
        rows.append([box])
    rows = [sorted(row, key=lambda b: b["box"][0]) for row in rows if len(row) >= 2]
    if len(rows) < 2:
        return []
    width = max(map(len, rows))
    exemplar = next(row for row in rows if len(row) == width)
    centers = [(b["box"][0] + b["box"][2]) / 2 for b in exemplar]
    if columns:
        centers = [(a + b) / 2 for a, b in zip(columns, columns[1:])]
        width = len(centers)
    matrix = []
    for row in rows:
        values = [None] * width
        for box in row:
            x = (box["box"][0] + box["box"][2]) / 2
            index = min(range(width), key=lambda i: abs(centers[i] - x))
            values[index] = (values[index] + " " if values[index] else "") + box["text"]
        matrix.append(values)
    bounds = [
        min(b["box"][0] for b in boxes),
        min(b["box"][1] for b in boxes),
        max(b["box"][2] for b in boxes),
        max(b["box"][3] for b in boxes),
    ]
    return [
        {
            "raw_rows": matrix,
            "bbox": bounds,
            "cells": [b["box"] for b in boxes],
            "engine": "rapidocr",
            "ocr_mean_confidence": sum(b["score"] for b in boxes) / len(boxes),
            "notes": [
                "OCR reconstruction is approximate; review row/column alignment and every number."
            ],
        }
    ]


def camelot_tables(
    path: Path,
    page: int,
    height: float,
    area: list[float] | None,
    columns: list[float] | None,
    flavor: str = "lattice",
) -> list[dict]:
    try:
        import camelot
    except ImportError as exc:
        raise PDFError(
            "missing_dependency",
            "Install Camelot support: pip install 'pdf-table-extract[camelot]'",
            3,
        ) from exc
    settings = {"pages": str(page), "flavor": flavor, "suppress_stdout": True}
    if area:
        settings["table_areas"] = [f"{area[0]},{height - area[1]},{area[2]},{height - area[3]}"]
    if columns and flavor == "stream":
        settings["columns"] = [",".join(map(str, columns[1:-1]))]
    try:
        extracted = camelot.read_pdf(str(path), **settings)
    except Exception as exc:
        raise PDFError("engine_failed", f"Camelot failed: {exc}") from exc
    tables = []
    for table in extracted:
        raw = table.df.values.tolist()
        if len(raw) < 2 or len(raw[0]) < 2:
            continue
        x0, y0, x1, y1 = table._bbox
        cells = [
            [cell.x1, height - cell.y2, cell.x2, height - cell.y1]
            for row in table.cells
            for cell in row
        ]
        tables.append(
            {
                "raw_rows": raw,
                "bbox": [x0, height - y1, x1, height - y0],
                "cells": cells,
                "engine": f"camelot-{flavor}",
                "notes": [
                    f"Camelot parsing accuracy: {table.parsing_report['accuracy']}% "
                    "(not a ground-truth guarantee)."
                ],
            }
        )
    return tables


def numeric(value: str | None) -> str | None:
    if value is None:
        return None
    text = value.strip().replace("−", "-")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    percentage = text.endswith("%")
    if percentage:
        text = text[:-1]
    if not re.fullmatch(r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?", text):
        return None
    clean = text.replace(",", "")
    if re.match(r"[+-]?0\d", clean):
        return None  # Preserve identifiers with significant leading zeroes.
    try:
        number = Decimal(clean)
        if negative:
            number = -number
        if percentage:
            number /= 100
        return format(number, "f")
    except InvalidOperation:
        return None


def clean_table(
    raw: dict,
    page: int,
    header_rows: int,
    fill_merged: str,
    normalize_numbers: bool,
    exclude_text: list[str],
) -> dict:
    rows = [
        [str(cell).replace("\n", " ").strip() if cell is not None else None for cell in row]
        for row in raw["raw_rows"]
    ]
    rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
    try:
        patterns = [re.compile(pattern) for pattern in exclude_text]
    except re.error as exc:
        raise PDFError("invalid_pattern", f"Invalid exclude-text expression: {exc}", 2) from exc
    rows = [
        row
        for row in rows
        if not any(pattern.search(" ".join(cell or "" for cell in row)) for pattern in patterns)
    ]
    if len(rows) <= header_rows:
        raise PDFError(
            "insufficient_rows",
            "Table has no data rows after the selected header/cleanup settings",
            2,
        )
    width = max(map(len, rows))
    rows = [row + [None] * (width - len(row)) for row in rows]
    headers_matrix = [list(row) for row in rows[:header_rows]]
    # Absent cells in merged upper headers inherit their parent label to the right.
    for row in headers_matrix:
        for index in range(1, width):
            if row[index] is None:
                row[index] = row[index - 1]
    headers = []
    for index in range(width):
        parts = list(dict.fromkeys(row[index] for row in headers_matrix if row[index]))
        header = "/".join(parts) or f"column_{index + 1}"
        headers.append(header)
    data = [list(row) for row in rows[header_rows:]]
    if fill_merged == "right":
        for row in data:
            for index in range(1, width):
                if row[index] is None:
                    row[index] = row[index - 1]
    elif fill_merged == "down":
        for row_index in range(1, len(data)):
            for index in range(width):
                if data[row_index][index] is None:
                    data[row_index][index] = data[row_index - 1][index]
    numeric_columns = []
    for column in range(width):
        values = [row[column] for row in data if row[column] not in (None, "")]
        if values and all(numeric(value) is not None for value in values):
            numeric_columns.append(column)
            if normalize_numbers:
                for row in data:
                    if row[column] not in (None, ""):
                        row[column] = numeric(row[column])
    empty = sum(cell in (None, "") for row in data for cell in row)
    notes = list(raw.get("notes", []))
    if empty:
        notes.append(
            "Empty cells may be intentional or may indicate extraction loss; inspect the overlay."
        )
    return {
        "id": "",
        "pages": [page],
        "headers": headers,
        "rows": data,
        "raw_rows": rows,
        "bbox": raw["bbox"],
        "engine": raw["engine"],
        "page_regions": [{"page": page, "bbox": raw["bbox"], "cells": raw["cells"]}],
        "quality": {
            "row_count": len(data),
            "column_count": width,
            "empty_cell_ratio": empty / (len(data) * width),
            "numeric_columns": numeric_columns,
            "notes": notes,
            "ocr_mean_confidence": raw.get("ocr_mean_confidence"),
        },
    }


def merge_contiguous(tables: list[dict]) -> list[dict]:
    result = []
    for table in tables:
        match = None
        for candidate in reversed(result):
            if (
                candidate["pages"][-1] + 1 == table["pages"][0]
                and len(candidate["headers"]) == len(table["headers"])
                and abs(candidate["bbox"][0] - table["bbox"][0]) <= 6
                and abs(candidate["bbox"][2] - table["bbox"][2]) <= 6
                and SequenceMatcher(
                    None,
                    "|".join(candidate["headers"]).casefold(),
                    "|".join(table["headers"]).casefold(),
                ).ratio()
                >= 0.85
            ):
                match = candidate
                break
        if match is None:
            result.append(table)
        else:
            match["rows"].extend(table["rows"])
            match["pages"].extend(table["pages"])
            match["page_regions"].extend(table["page_regions"])
            match["raw_rows"].extend(table["raw_rows"])
            match["quality"]["row_count"] = len(match["rows"])
            total = len(match["rows"]) * len(match["headers"])
            match["quality"]["empty_cell_ratio"] = (
                sum(c in (None, "") for r in match["rows"] for c in r) / total
            )
            match["quality"]["notes"].append(
                "Merged consecutive pages using matching headers and horizontal geometry."
            )
    return result


def inspect_pdf(path: str | Path) -> dict:
    with open_pdf(path) as document:
        pages = []
        for index, page in enumerate(document.pages, 1):
            has_text = bool(page.chars)
            tables = plumber_tables(page) if has_text else []
            pages.append(
                {
                    "page": index,
                    "width": float(page.width),
                    "height": float(page.height),
                    "has_text": has_text,
                    "characters": len(page.chars),
                    "table_count": len(tables),
                    "image_count": len(page.images),
                }
            )
        return {"name": Path(path).name, "page_count": len(pages), "pages": pages}


def extract_tables(
    path: str | Path,
    *,
    pages: str | None = None,
    out: str | Path | None = None,
    format: str = "json",
    engine: str = "auto",
    ocr: bool = False,
    merge_pages: bool = False,
    fill_merged: str = "none",
    header_rows: int = 1,
    area: list[float] | None = None,
    columns: list[float] | None = None,
    normalize_numbers: bool = True,
    trim_top: float = 0,
    trim_bottom: float = 0,
    exclude_text: list[str] | None = None,
    compare: bool = False,
    camelot_flavor: str = "lattice",
    bom: bool = True,
    progress: Callable[[str], None] | None = None,
) -> dict:
    if engine not in {"auto", "plumber", "camelot"} or fill_merged not in {"none", "down", "right"}:
        raise PDFError("invalid_option", "Unknown engine or fill-merged option", 2)
    if format not in {"csv", "xlsx", "json"}:
        raise PDFError("invalid_format", "format must be csv, xlsx, or json", 2)
    if not 1 <= header_rows <= 20 or trim_top < 0 or trim_bottom < 0:
        raise PDFError("invalid_option", "header_rows must be 1..20 and trims nonnegative", 2)
    if camelot_flavor not in {"lattice", "stream"}:
        raise PDFError("invalid_option", "Camelot flavor must be lattice or stream", 2)
    path = Path(path).expanduser()
    tables, warnings, comparisons = [], [], []
    with open_pdf(path) as document:
        count = len(document.pages)
        for number in selected_pages(pages, count):
            if progress:
                progress(f"Extracting page {number}/{count}")
            original = document.pages[number - 1]
            page = page_area(original, area, trim_top, trim_bottom)
            camelot_area = list(page.bbox) if area or trim_top or trim_bottom else None
            if not original.chars:
                if not ocr:
                    warnings.append(f"Page {number} has no text layer; use --ocr for local OCR.")
                    continue
                found = ocr_tables(path, number, list(page.bbox), columns)
            elif engine == "camelot":
                found = camelot_tables(
                    path, number, original.height, camelot_area, columns, camelot_flavor
                )
            else:
                found = plumber_tables(page, columns)
            for raw in sorted(found, key=lambda table: (table["bbox"][1], table["bbox"][0])):
                tables.append(
                    clean_table(
                        raw, number, header_rows, fill_merged, normalize_numbers, exclude_text or []
                    )
                )
            if compare and original.chars:
                alternate = camelot_tables(
                    path, number, original.height, camelot_area, columns, camelot_flavor
                )
                baseline = plumber_tables(page, columns)
                comparisons.append(
                    {
                        "page": number,
                        "plumber": [t["raw_rows"] for t in baseline],
                        "camelot": [t["raw_rows"] for t in alternate],
                        "same_raw_tables": [t["raw_rows"] for t in baseline]
                        == [t["raw_rows"] for t in alternate],
                    }
                )
    if merge_pages:
        tables = merge_contiguous(tables)
    for index, table in enumerate(tables, 1):
        table["id"] = f"table_{index:03d}"
    if not tables:
        warnings.append(
            "No tables extracted. Review page selection, regions, strategy, or OCR support."
        )
    result = {
        "schema_version": 1,
        "source": {
            "name": path.name,
            "sha1": hashlib.sha1(path.read_bytes(), usedforsecurity=False).hexdigest(),
            "page_count": count,
        },
        "tables": tables,
        "warnings": warnings,
        "comparisons": comparisons,
        "files": [],
    }
    errors = list(RESULT_VALIDATOR.iter_errors(result))
    if errors:
        raise PDFError("invalid_result", errors[0].message)
    if out:
        result["files"] = write_results(result, Path(out), format, bom)
    return result


def write_results(result: dict, directory: Path, format: str, bom: bool = True) -> list[str]:
    if format not in {"json", "csv", "xlsx"}:
        raise PDFError("invalid_format", "format must be json, csv, or xlsx", 2)
    directory.mkdir(parents=True, exist_ok=True)
    files = []
    if format == "json":
        path = directory / "tables.json"
        path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        files.append(str(path.resolve()))
    elif format == "csv":
        for table in result["tables"]:
            path = directory / f"{table['id']}.csv"
            with path.open("w", encoding="utf-8-sig" if bom else "utf-8", newline="") as file:
                csv.writer(file).writerows([table["headers"], *table["rows"]])
            files.append(str(path.resolve()))
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for table in result["tables"]:
            sheet = workbook.create_sheet(table["id"])
            for row in [table["headers"], *table["rows"]]:
                sheet.append(row)
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        cell.data_type = "s"  # Preserve decimal text and prevent formula execution.
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        if not result["tables"]:
            workbook.create_sheet("No tables")
        path = directory / "tables.xlsx"
        workbook.save(path)
        files.append(str(path.resolve()))
    quality = directory / "quality.json"
    quality.write_text(
        json.dumps(
            {
                "source": result["source"],
                "warnings": result["warnings"],
                "tables": [{"id": t["id"], **t["quality"]} for t in result["tables"]],
            },
            indent=2,
        )
        + "\n"
    )
    files.append(str(quality.resolve()))
    return files
