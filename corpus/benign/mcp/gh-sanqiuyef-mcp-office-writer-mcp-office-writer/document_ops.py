"""Path-addressed inspection and transactional editing for OOXML documents."""
from __future__ import annotations

import os
import re
from copy import copy
from pathlib import Path
from typing import Any, Iterable

from pydantic import ValidationError

from .models import OfficeOperation
from .utils import atomic_save, ensure_output_dir, validate_input_file, validate_output_path


class OfficeOperationError(Exception):
    """Error payload designed for an agent to correct and retry."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        path: str | None = None,
        suggestion: str | None = None,
        valid_values: Iterable[str] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.path = path
        self.suggestion = suggestion
        self.valid_values = list(valid_values) if valid_values is not None else None

    def to_dict(self) -> dict[str, Any]:
        result: dict[str, Any] = {"code": self.code, "message": str(self)}
        if self.path is not None:
            result["path"] = self.path
        if self.suggestion is not None:
            result["suggestion"] = self.suggestion
        if self.valid_values is not None:
            result["valid_values"] = self.valid_values
        return result


class _NodeBudget:
    def __init__(self, maximum: int) -> None:
        self.maximum = maximum
        self.count = 0
        self.truncated = False

    def add(self, parent: dict[str, Any], node: dict[str, Any]) -> bool:
        if self.count >= self.maximum:
            self.truncated = True
            return False
        parent.setdefault("children", []).append(node)
        self.count += 1
        return True


def _format_from_suffix(file_path: str) -> str:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".docx":
        return "docx"
    if suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    if suffix == ".pptx":
        return "pptx"
    raise OfficeOperationError(
        "unsupported_type",
        f"不支持的文件格式: {suffix or '(无扩展名)'}",
        path=file_path,
        valid_values=[".docx", ".xlsx", ".xlsm", ".pptx"],
    )


def _node(path: str, node_type: str, *, text: str | None = None, **values: Any) -> dict[str, Any]:
    result: dict[str, Any] = {"path": path, "type": node_type, "children": []}
    if text is not None:
        result["text"] = text
    result.update(values)
    return result


def _inspect_word(file_path: str, budget: _NodeBudget) -> dict[str, Any]:
    from docx import Document

    doc = Document(file_path)
    root = _node("/", "document", format="docx")
    budget.count = 1
    body = _node("/body", "body")
    if not budget.add(root, body):
        return root

    for index, paragraph in enumerate(doc.paragraphs, 1):
        style_name = paragraph.style.name if paragraph.style else None
        paragraph_node = _node(
            f"/body/p[{index}]",
            "paragraph",
            text=paragraph.text,
            style=style_name,
            format={"alignment": str(paragraph.alignment) if paragraph.alignment is not None else None},
        )
        if not budget.add(body, paragraph_node):
            break

    for table_index, table in enumerate(doc.tables, 1):
        table_node = _node(
            f"/body/tbl[{table_index}]",
            "table",
            format={"rows": len(table.rows), "columns": len(table.columns)},
        )
        if not budget.add(body, table_node):
            break
        for row_index, row in enumerate(table.rows, 1):
            row_node = _node(f"/body/tbl[{table_index}]/row[{row_index}]", "row")
            if not budget.add(table_node, row_node):
                break
            for cell_index, cell in enumerate(row.cells, 1):
                cell_node = _node(
                    f"/body/tbl[{table_index}]/row[{row_index}]/cell[{cell_index}]",
                    "cell",
                    text=cell.text,
                )
                if not budget.add(row_node, cell_node):
                    break
    return root


def _inspect_excel(file_path: str, budget: _NodeBudget) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=False, data_only=False, keep_vba=file_path.lower().endswith(".xlsm"))
    try:
        root = _node("/", "workbook", format="xlsx")
        budget.count = 1
        for ws in wb.worksheets:
            sheet_node = _node(
                f"/{ws.title}",
                "sheet",
                format={"max_row": ws.max_row, "max_column": ws.max_column, "state": ws.sheet_state},
            )
            if not budget.add(root, sheet_node):
                break
            populated = sorted(ws._cells.values(), key=lambda cell: (cell.row, cell.column))
            for cell in populated:
                if cell.value is None:
                    continue
                cell_node = _node(
                    f"/{ws.title}/{cell.coordinate}",
                    "cell",
                    text=str(cell.value),
                    value=cell.value,
                    format={
                        "number_format": cell.number_format,
                        "bold": bool(cell.font.bold),
                        "fill": cell.fill.fgColor.rgb if cell.fill.fill_type else None,
                    },
                )
                if not budget.add(sheet_node, cell_node):
                    break
        return root
    finally:
        wb.close()


def _shape_text(shape: Any) -> str | None:
    if getattr(shape, "has_text_frame", False):
        return shape.text
    if getattr(shape, "has_table", False):
        return "\n".join(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
    return None


def _inspect_powerpoint(file_path: str, budget: _NodeBudget) -> dict[str, Any]:
    from pptx import Presentation

    prs = Presentation(file_path)
    root = _node("/", "presentation", format="pptx")
    budget.count = 1
    for slide_index, slide in enumerate(prs.slides, 1):
        slide_node = _node(
            f"/slide[{slide_index}]",
            "slide",
            format={"layout": slide.slide_layout.name if slide.slide_layout else None},
        )
        if not budget.add(root, slide_node):
            break
        for shape_index, shape in enumerate(slide.shapes, 1):
            shape_node = _node(
                f"/slide[{slide_index}]/shape[{shape_index}]",
                "shape",
                text=_shape_text(shape),
                name=shape.name,
                format={
                    "shape_type": str(shape.shape_type),
                    "x": round(shape.left / 914400, 3),
                    "y": round(shape.top / 914400, 3),
                    "width": round(shape.width / 914400, 3),
                    "height": round(shape.height / 914400, 3),
                },
            )
            if not budget.add(slide_node, shape_node):
                break
    return root


def inspect_office_document(file_path: str, max_nodes: int = 1000) -> dict[str, Any]:
    """Return an agent-addressable node tree for a Word, Excel, or PowerPoint file."""
    if max_nodes < 1 or max_nodes > 10000:
        return {
            "success": False,
            "file_path": file_path,
            "error": OfficeOperationError(
                "invalid_value", "max_nodes 必须在 1 到 10000 之间", valid_values=["1..10000"]
            ).to_dict(),
        }
    try:
        file_format = _format_from_suffix(file_path)
        allowed = {"docx": (".docx",), "xlsx": (".xlsx", ".xlsm"), "pptx": (".pptx",)}[file_format]
        validate_input_file(file_path, allowed)
        budget = _NodeBudget(max_nodes)
        if file_format == "docx":
            root = _inspect_word(file_path, budget)
        elif file_format == "xlsx":
            root = _inspect_excel(file_path, budget)
        else:
            root = _inspect_powerpoint(file_path, budget)
        return {
            "success": True,
            "file_path": file_path,
            "format": file_format,
            "node_count": budget.count,
            "truncated": budget.truncated,
            "root": root,
        }
    except OfficeOperationError as exc:
        return {"success": False, "file_path": file_path, "error": exc.to_dict()}
    except ValueError as exc:
        return {
            "success": False,
            "file_path": file_path,
            "error": OfficeOperationError("invalid_input", str(exc), path=file_path).to_dict(),
        }
    except Exception as exc:
        return {
            "success": False,
            "file_path": file_path,
            "error": OfficeOperationError("read_failed", f"读取文档结构失败: {exc}", path=file_path).to_dict(),
        }


_WORD_PARAGRAPH = re.compile(r"^/body/p\[(\d+)\]$")
_WORD_TABLE = re.compile(r"^/body/tbl\[(\d+)\]$")
_WORD_CELL = re.compile(r"^/body/tbl\[(\d+)\]/row\[(\d+)\]/cell\[(\d+)\]$")


def _one_based(items: Any, index: int, path: str) -> Any:
    if index < 1 or index > len(items):
        raise OfficeOperationError(
            "not_found", f"路径索引超出范围: {path}", path=path, suggestion=f"有效索引为 1..{len(items)}"
        )
    return items[index - 1]


def _word_target(doc: Any, path: str) -> tuple[str, Any]:
    match = _WORD_PARAGRAPH.fullmatch(path)
    if match:
        return "paragraph", _one_based(doc.paragraphs, int(match.group(1)), path)
    match = _WORD_TABLE.fullmatch(path)
    if match:
        return "table", _one_based(doc.tables, int(match.group(1)), path)
    match = _WORD_CELL.fullmatch(path)
    if match:
        table = _one_based(doc.tables, int(match.group(1)), path)
        row = _one_based(table.rows, int(match.group(2)), path)
        return "cell", _one_based(row.cells, int(match.group(3)), path)
    if path == "/body":
        return "body", doc
    raise OfficeOperationError(
        "invalid_path",
        f"不支持的 Word 路径: {path}",
        path=path,
        suggestion="使用 inspect_office_document 获取路径",
        valid_values=["/body", "/body/p[N]", "/body/tbl[N]", "/body/tbl[N]/row[N]/cell[N]"],
    )


def _replace_paragraph_text(paragraph: Any, find: str, replacement: str) -> int:
    original = paragraph.text
    count = original.count(find)
    if not count:
        return 0
    changed = original.replace(find, replacement)
    if paragraph.runs:
        paragraph.runs[0].text = changed
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(changed)
    return count


def _word_paragraphs(doc: Any) -> Iterable[Any]:
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _set_word_paragraph(paragraph: Any, props: dict[str, Any], path: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    allowed = {"text", "style", "bold", "italic", "font_name", "font_size", "font_color", "alignment"}
    unknown = set(props) - allowed
    if unknown:
        raise OfficeOperationError(
            "unsupported_property", f"Word 段落不支持属性: {sorted(unknown)}", path=path, valid_values=sorted(allowed)
        )
    if "text" in props:
        if paragraph.runs:
            paragraph.runs[0].text = str(props["text"])
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(str(props["text"]))
    if "style" in props:
        paragraph.style = str(props["style"])
    if "alignment" in props:
        alignments = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        value = str(props["alignment"]).lower()
        if value not in alignments:
            raise OfficeOperationError("invalid_value", f"无效的对齐方式: {value}", path=path, valid_values=alignments)
        paragraph.alignment = alignments[value]
    for run in paragraph.runs:
        if "bold" in props:
            run.bold = bool(props["bold"])
        if "italic" in props:
            run.italic = bool(props["italic"])
        if "font_name" in props:
            run.font.name = str(props["font_name"])
        if "font_size" in props:
            run.font.size = Pt(float(props["font_size"]))
        if "font_color" in props:
            run.font.color.rgb = RGBColor.from_string(str(props["font_color"]).lstrip("#"))


def _apply_word_operation(doc: Any, operation: OfficeOperation) -> str:
    kind, target = _word_target(doc, operation.path)
    if operation.op == "set":
        if kind == "paragraph":
            _set_word_paragraph(target, operation.props, operation.path)
        elif kind == "cell":
            unknown = set(operation.props) - {"text"}
            if unknown:
                raise OfficeOperationError(
                    "unsupported_property", f"Word 单元格不支持属性: {sorted(unknown)}", path=operation.path, valid_values=["text"]
                )
            if "text" in operation.props:
                target.text = str(operation.props["text"])
        else:
            raise OfficeOperationError("unsupported_operation", f"set 不支持目标类型: {kind}", path=operation.path)
        return f"已更新 {operation.path}"

    if operation.op == "replace":
        find = str(operation.props.get("find", ""))
        replacement = str(operation.props.get("replace", ""))
        if not find:
            raise OfficeOperationError("invalid_value", "replace 操作的 find 不能为空", path=operation.path)
        paragraphs: Iterable[Any]
        if kind == "body":
            paragraphs = _word_paragraphs(doc)
        elif kind == "paragraph":
            paragraphs = [target]
        elif kind == "cell":
            paragraphs = target.paragraphs
        else:
            raise OfficeOperationError("unsupported_operation", "replace 仅支持正文、段落或单元格", path=operation.path)
        count = sum(_replace_paragraph_text(paragraph, find, replacement) for paragraph in paragraphs)
        return f"已替换 {count} 处文本"

    if operation.op == "add":
        if kind != "body":
            raise OfficeOperationError("unsupported_operation", "Word add 当前仅支持 /body", path=operation.path)
        element_type = operation.element_type
        if element_type == "paragraph":
            paragraph = doc.add_paragraph(str(operation.props.get("text", "")))
            if "style" in operation.props:
                paragraph.style = str(operation.props["style"])
        elif element_type == "heading":
            doc.add_heading(str(operation.props.get("text", "")), level=int(operation.props.get("level", 1)))
        elif element_type == "page_break":
            doc.add_page_break()
        elif element_type == "table":
            rows = operation.props.get("rows", [])
            headers = operation.props.get("headers", [])
            column_count = max(len(headers), max((len(row) for row in rows), default=0))
            if column_count < 1:
                raise OfficeOperationError("invalid_value", "table 至少需要一列", path=operation.path)
            table = doc.add_table(rows=len(rows) + (1 if headers else 0), cols=column_count)
            table.style = "Table Grid"
            all_rows = ([headers] if headers else []) + list(rows)
            for row_index, values in enumerate(all_rows):
                for column_index, value in enumerate(values):
                    table.cell(row_index, column_index).text = str(value)
        else:
            raise OfficeOperationError(
                "invalid_value",
                f"不支持的 Word 元素类型: {element_type}",
                path=operation.path,
                valid_values=["paragraph", "heading", "page_break", "table"],
            )
        return f"已添加 {element_type}"

    if operation.op == "remove":
        if kind in {"paragraph", "table"}:
            element = target._element
            element.getparent().remove(element)
        elif kind == "cell":
            target.text = ""
        else:
            raise OfficeOperationError("unsupported_operation", "不能删除整个 Word 正文", path=operation.path)
        return f"已删除 {operation.path}"

    if operation.op == "move":
        if kind not in {"paragraph", "table"} or not operation.target_path:
            raise OfficeOperationError("invalid_value", "Word move 需要正文段落/表格和 target_path", path=operation.path)
        target_kind, move_target = _word_target(doc, operation.target_path)
        if target_kind not in {"paragraph", "table"}:
            raise OfficeOperationError("invalid_path", "target_path 必须是正文段落或表格", path=operation.target_path)
        if target is move_target:
            raise OfficeOperationError("invalid_value", "源路径和目标路径不能相同", path=operation.path)
        if operation.position == "before":
            move_target._element.addprevious(target._element)
        else:
            move_target._element.addnext(target._element)
        return f"已移动 {operation.path}"

    raise OfficeOperationError("unsupported_operation", f"不支持的操作: {operation.op}", path=operation.path)


_EXCEL_PATH = re.compile(r"^/([^/]+)(?:/([A-Za-z]+\d+(?::[A-Za-z]+\d+)?|row\[\d+\]|col\[[A-Za-z]+\]))?$")


def _excel_target(wb: Any, path: str) -> tuple[Any | None, str | None]:
    if path == "/":
        return None, None
    match = _EXCEL_PATH.fullmatch(path)
    if not match:
        raise OfficeOperationError(
            "invalid_path",
            f"不支持的 Excel 路径: {path}",
            path=path,
            valid_values=["/", "/SheetName", "/SheetName/A1", "/SheetName/A1:C3", "/SheetName/row[N]", "/SheetName/col[A]"],
        )
    name, selector = match.groups()
    if name not in wb.sheetnames:
        raise OfficeOperationError("not_found", f"工作表不存在: {name}", path=path, valid_values=wb.sheetnames)
    return wb[name], selector


def _excel_cells(ws: Any, selector: str | None) -> list[Any]:
    from openpyxl.utils.cell import range_boundaries

    if selector is None:
        return [cell for cell in ws._cells.values()]
    if selector.startswith("row[") or selector.startswith("col["):
        raise OfficeOperationError("invalid_path", "此操作需要单元格或区域路径")
    min_col, min_row, max_col, max_row = range_boundaries(selector)
    return [ws.cell(row=row, column=column) for row in range(min_row, max_row + 1) for column in range(min_col, max_col + 1)]


def _set_excel_cell(cell: Any, props: dict[str, Any], path: str) -> None:
    from openpyxl.styles import PatternFill

    allowed = {"value", "formula", "bold", "italic", "font_name", "font_size", "font_color", "fill", "alignment", "number_format"}
    unknown = set(props) - allowed
    if unknown:
        raise OfficeOperationError(
            "unsupported_property", f"Excel 单元格不支持属性: {sorted(unknown)}", path=path, valid_values=sorted(allowed)
        )
    if "value" in props:
        value = props["value"]
        cell.value = "'" + value if isinstance(value, str) and value.startswith("=") else value
    if "formula" in props:
        formula = str(props["formula"])
        cell.value = formula if formula.startswith("=") else "=" + formula
    font_updates: dict[str, Any] = {}
    for key, field in (("bold", "bold"), ("italic", "italic"), ("font_name", "name"), ("font_size", "size")):
        if key in props:
            font_updates[field] = props[key]
    if "font_color" in props:
        font_updates["color"] = str(props["font_color"]).lstrip("#")
    if font_updates:
        updated_font = copy(cell.font)
        for field, value in font_updates.items():
            setattr(updated_font, field, value)
        cell.font = updated_font
    if "fill" in props:
        cell.fill = PatternFill("solid", fgColor=str(props["fill"]).lstrip("#"))
    if "alignment" in props:
        value = str(props["alignment"]).lower()
        if value not in {"left", "center", "right"}:
            raise OfficeOperationError("invalid_value", f"无效的对齐方式: {value}", path=path, valid_values=["left", "center", "right"])
        updated_alignment = copy(cell.alignment)
        updated_alignment.horizontal = value
        cell.alignment = updated_alignment
    if "number_format" in props:
        cell.number_format = str(props["number_format"])


def _apply_excel_operation(wb: Any, operation: OfficeOperation) -> str:
    ws, selector = _excel_target(wb, operation.path)
    if operation.op == "set":
        if ws is None or selector is None or selector.startswith(("row[", "col[")):
            raise OfficeOperationError("invalid_path", "Excel set 需要单元格或区域路径", path=operation.path)
        cells = _excel_cells(ws, selector)
        if "values" in operation.props:
            values = operation.props["values"]
            from openpyxl.utils.cell import range_boundaries
            min_col, min_row, max_col, max_row = range_boundaries(selector)
            expected_rows, expected_columns = max_row - min_row + 1, max_col - min_col + 1
            if len(values) != expected_rows or any(len(row) != expected_columns for row in values):
                raise OfficeOperationError(
                    "invalid_value", f"values 尺寸必须为 {expected_rows}x{expected_columns}", path=operation.path
                )
            for row_index, row in enumerate(values):
                for column_index, value in enumerate(row):
                    ws.cell(min_row + row_index, min_col + column_index).value = value
            props = {key: value for key, value in operation.props.items() if key != "values"}
        else:
            props = operation.props
        for cell in cells:
            _set_excel_cell(cell, props, operation.path)
        return f"已更新 {len(cells)} 个单元格"

    if operation.op == "replace":
        if ws is None:
            raise OfficeOperationError("invalid_path", "Excel replace 需要工作表或区域路径", path=operation.path)
        find = str(operation.props.get("find", ""))
        replacement = str(operation.props.get("replace", ""))
        if not find:
            raise OfficeOperationError("invalid_value", "replace 操作的 find 不能为空", path=operation.path)
        cells = _excel_cells(ws, selector)
        count = 0
        for cell in cells:
            if isinstance(cell.value, str) and find in cell.value:
                count += cell.value.count(find)
                cell.value = cell.value.replace(find, replacement)
        return f"已替换 {count} 处文本"

    if operation.op == "add":
        if ws is None:
            if operation.element_type != "sheet":
                raise OfficeOperationError("invalid_value", "工作簿根节点只能添加 sheet", path=operation.path, valid_values=["sheet"])
            name = str(operation.props.get("name", "Sheet"))[:31]
            if name in wb.sheetnames:
                raise OfficeOperationError("already_exists", f"工作表已存在: {name}", path=operation.path)
            wb.create_sheet(name)
            return f"已添加工作表 {name}"
        if selector is not None:
            raise OfficeOperationError("invalid_path", "添加行或列时路径应指向工作表", path=operation.path)
        index = int(operation.props.get("index", ws.max_row + 1 if operation.element_type == "row" else ws.max_column + 1))
        values = operation.props.get("values", [])
        if operation.element_type == "row":
            ws.insert_rows(index)
            for column, value in enumerate(values, 1):
                ws.cell(index, column).value = value
        elif operation.element_type == "column":
            ws.insert_cols(index)
            for row, value in enumerate(values, 1):
                ws.cell(row, index).value = value
        else:
            raise OfficeOperationError("invalid_value", f"不支持的 Excel 元素类型: {operation.element_type}", valid_values=["row", "column"])
        return f"已添加 {operation.element_type}"

    if operation.op == "remove":
        if ws is None:
            raise OfficeOperationError("unsupported_operation", "不能删除整个工作簿", path=operation.path)
        if selector is None:
            if len(wb.worksheets) == 1:
                raise OfficeOperationError("invalid_value", "工作簿必须至少保留一个工作表", path=operation.path)
            wb.remove(ws)
        elif selector.startswith("row["):
            ws.delete_rows(int(selector[4:-1]))
        elif selector.startswith("col["):
            from openpyxl.utils.cell import column_index_from_string
            ws.delete_cols(column_index_from_string(selector[4:-1]))
        else:
            for cell in _excel_cells(ws, selector):
                cell.value = None
        return f"已删除 {operation.path}"

    if operation.op == "move":
        if ws is None or selector is not None or not operation.target_path:
            raise OfficeOperationError("unsupported_operation", "Excel move 当前仅支持工作表排序", path=operation.path)
        target_ws, target_selector = _excel_target(wb, operation.target_path)
        if target_ws is None or target_selector is not None or target_ws is ws:
            raise OfficeOperationError("invalid_path", "target_path 必须是另一个工作表", path=operation.target_path)
        wb._sheets.remove(ws)
        target_index = wb._sheets.index(target_ws)
        wb._sheets.insert(target_index if operation.position == "before" else target_index + 1, ws)
        return f"已移动工作表 {ws.title}"

    raise OfficeOperationError("unsupported_operation", f"不支持的操作: {operation.op}", path=operation.path)


_PPT_SLIDE = re.compile(r"^/slide\[(\d+)\]$")
_PPT_SHAPE = re.compile(r"^/slide\[(\d+)\]/shape\[(\d+)\]$")


def _ppt_target(prs: Any, path: str) -> tuple[str, Any]:
    if path == "/":
        return "presentation", prs
    match = _PPT_SHAPE.fullmatch(path)
    if match:
        slide = _one_based(prs.slides, int(match.group(1)), path)
        return "shape", _one_based(slide.shapes, int(match.group(2)), path)
    match = _PPT_SLIDE.fullmatch(path)
    if match:
        return "slide", _one_based(prs.slides, int(match.group(1)), path)
    raise OfficeOperationError(
        "invalid_path", f"不支持的 PPT 路径: {path}", path=path, valid_values=["/", "/slide[N]", "/slide[N]/shape[N]"]
    )


def _set_ppt_shape(shape: Any, props: dict[str, Any], path: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    allowed = {"text", "x", "y", "width", "height", "fill", "font_color", "font_size", "bold", "italic"}
    unknown = set(props) - allowed
    if unknown:
        raise OfficeOperationError(
            "unsupported_property", f"PPT 形状不支持属性: {sorted(unknown)}", path=path, valid_values=sorted(allowed)
        )
    if "text" in props:
        if not getattr(shape, "has_text_frame", False):
            raise OfficeOperationError("unsupported_operation", "目标形状不包含文本框", path=path)
        shape.text = str(props["text"])
    for key, attr in (("x", "left"), ("y", "top"), ("width", "width"), ("height", "height")):
        if key in props:
            setattr(shape, attr, Inches(float(props[key])))
    if "fill" in props:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(str(props["fill"]).lstrip("#"))
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if "font_color" in props:
                    run.font.color.rgb = RGBColor.from_string(str(props["font_color"]).lstrip("#"))
                if "font_size" in props:
                    run.font.size = Pt(float(props["font_size"]))
                if "bold" in props:
                    run.font.bold = bool(props["bold"])
                if "italic" in props:
                    run.font.italic = bool(props["italic"])


def _apply_ppt_operation(prs: Any, operation: OfficeOperation) -> str:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    kind, target = _ppt_target(prs, operation.path)
    if operation.op == "set":
        if kind == "shape":
            _set_ppt_shape(target, operation.props, operation.path)
        elif kind == "slide":
            allowed = {"title", "notes", "background"}
            unknown = set(operation.props) - allowed
            if unknown:
                raise OfficeOperationError("unsupported_property", f"PPT 幻灯片不支持属性: {sorted(unknown)}", valid_values=sorted(allowed))
            if "title" in operation.props:
                if target.shapes.title is None:
                    raise OfficeOperationError("not_found", "幻灯片没有标题占位符", path=operation.path)
                target.shapes.title.text = str(operation.props["title"])
            if "notes" in operation.props:
                target.notes_slide.notes_text_frame.text = str(operation.props["notes"])
            if "background" in operation.props:
                target.background.fill.solid()
                target.background.fill.fore_color.rgb = RGBColor.from_string(str(operation.props["background"]).lstrip("#"))
        else:
            raise OfficeOperationError("unsupported_operation", "set 不支持整个演示文稿", path=operation.path)
        return f"已更新 {operation.path}"

    if operation.op == "replace":
        find = str(operation.props.get("find", ""))
        replacement = str(operation.props.get("replace", ""))
        if not find:
            raise OfficeOperationError("invalid_value", "replace 操作的 find 不能为空", path=operation.path)
        shapes: list[Any]
        if kind == "presentation":
            shapes = [shape for slide in prs.slides for shape in slide.shapes]
        elif kind == "slide":
            shapes = list(target.shapes)
        else:
            shapes = [target]
        count = 0
        for shape in shapes:
            if getattr(shape, "has_text_frame", False) and find in shape.text:
                count += shape.text.count(find)
                shape.text = shape.text.replace(find, replacement)
        return f"已替换 {count} 处文本"

    if operation.op == "add":
        if kind == "presentation":
            if operation.element_type != "slide":
                raise OfficeOperationError("invalid_value", "演示文稿根节点只能添加 slide", valid_values=["slide"])
            layout_index = int(operation.props.get("layout_index", 6))
            if layout_index < 0 or layout_index >= len(prs.slide_layouts):
                raise OfficeOperationError("invalid_value", f"layout_index 超出范围: {layout_index}")
            slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
            title = operation.props.get("title")
            if title is not None and slide.shapes.title is not None:
                slide.shapes.title.text = str(title)
            return "已添加 slide"
        if kind != "slide":
            raise OfficeOperationError("invalid_path", "PPT 元素应添加到幻灯片路径", path=operation.path)
        x = Inches(float(operation.props.get("x", 1)))
        y = Inches(float(operation.props.get("y", 1)))
        width = Inches(float(operation.props.get("width", 4)))
        height = Inches(float(operation.props.get("height", 1)))
        if operation.element_type == "textbox":
            shape = target.shapes.add_textbox(x, y, width, height)
            shape.text = str(operation.props.get("text", ""))
        elif operation.element_type == "shape":
            shape_types = {
                "rect": MSO_SHAPE.RECTANGLE,
                "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
                "oval": MSO_SHAPE.OVAL,
                "arrow": MSO_SHAPE.RIGHT_ARROW,
            }
            shape_name = str(operation.props.get("shape_type", "rect"))
            if shape_name not in shape_types:
                raise OfficeOperationError("invalid_value", f"无效的 shape_type: {shape_name}", valid_values=shape_types)
            shape = target.shapes.add_shape(shape_types[shape_name], x, y, width, height)
            shape.text = str(operation.props.get("text", ""))
        else:
            raise OfficeOperationError("invalid_value", f"不支持的 PPT 元素类型: {operation.element_type}", valid_values=["textbox", "shape"])
        return f"已添加 {operation.element_type}"

    if operation.op == "remove":
        if kind == "shape":
            element = target._element
            element.getparent().remove(element)
        elif kind == "slide":
            slide_id = target.slide_id
            slide_ids = prs.slides._sldIdLst
            slide_id_element = next(item for item in slide_ids if int(item.id) == slide_id)
            relationship_id = slide_id_element.rId
            prs.part.drop_rel(relationship_id)
            slide_ids.remove(slide_id_element)
        else:
            raise OfficeOperationError("unsupported_operation", "不能删除整个演示文稿", path=operation.path)
        return f"已删除 {operation.path}"

    if operation.op == "move":
        if not operation.target_path:
            raise OfficeOperationError("invalid_value", "move 操作需要 target_path", path=operation.path)
        target_kind, move_target = _ppt_target(prs, operation.target_path)
        if target_kind != kind or target is move_target:
            raise OfficeOperationError("invalid_path", "源路径与 target_path 必须是不同的同类元素", path=operation.target_path)
        if kind == "shape":
            if operation.position == "before":
                move_target._element.addprevious(target._element)
            else:
                move_target._element.addnext(target._element)
        elif kind == "slide":
            slide_ids = prs.slides._sldIdLst
            source_element = next(item for item in slide_ids if int(item.id) == target.slide_id)
            target_element = next(item for item in slide_ids if int(item.id) == move_target.slide_id)
            slide_ids.remove(source_element)
            target_index = list(slide_ids).index(target_element)
            slide_ids.insert(target_index if operation.position == "before" else target_index + 1, source_element)
        else:
            raise OfficeOperationError("unsupported_operation", "只能移动幻灯片或形状", path=operation.path)
        return f"已移动 {operation.path}"

    raise OfficeOperationError("unsupported_operation", f"不支持的操作: {operation.op}", path=operation.path)


def _parse_operations(operations: list[OfficeOperation | dict[str, Any]]) -> list[OfficeOperation]:
    parsed: list[OfficeOperation] = []
    for index, operation in enumerate(operations):
        try:
            parsed.append(operation if isinstance(operation, OfficeOperation) else OfficeOperation.model_validate(operation))
        except ValidationError as exc:
            raise OfficeOperationError(
                "invalid_operation", f"第 {index + 1} 个操作参数无效: {exc.errors(include_url=False)}"
            ) from exc
    if not parsed:
        raise OfficeOperationError("invalid_operation", "operations 不能为空")
    return parsed


def batch_edit_office_document(
    file_path: str,
    operations: list[OfficeOperation | dict[str, Any]],
    output_path: str | None = None,
) -> dict[str, Any]:
    """Apply path-addressed operations as one all-or-nothing atomic transaction."""
    results: list[dict[str, Any]] = []
    workbook = None
    try:
        file_format = _format_from_suffix(file_path)
        allowed = {"docx": (".docx",), "xlsx": (".xlsx", ".xlsm"), "pptx": (".pptx",)}[file_format]
        validate_input_file(file_path, allowed)
        target_path = output_path or file_path
        validate_output_path(target_path, allowed)
        if Path(target_path).suffix.lower() != Path(file_path).suffix.lower():
            raise OfficeOperationError(
                "extension_mismatch", "output_path 必须与源文件使用相同扩展名", path=target_path
            )
        ensure_output_dir(target_path)
        parsed = _parse_operations(operations)

        if file_format == "docx":
            from docx import Document
            document = Document(file_path)
            def apply(operation: OfficeOperation) -> str:
                return _apply_word_operation(document, operation)
            save = document.save
        elif file_format == "xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(file_path, keep_vba=file_path.lower().endswith(".xlsm"))
            def apply(operation: OfficeOperation) -> str:
                return _apply_excel_operation(workbook, operation)
            save = workbook.save
        else:
            from pptx import Presentation
            presentation = Presentation(file_path)
            def apply(operation: OfficeOperation) -> str:
                return _apply_ppt_operation(presentation, operation)
            save = presentation.save

        for index, operation in enumerate(parsed):
            try:
                message = apply(operation)
                results.append({"index": index, "success": True, "op": operation.op, "path": operation.path, "message": message})
            except OfficeOperationError as exc:
                results.append({"index": index, "success": False, "op": operation.op, "path": operation.path, "error": exc.to_dict()})
                return {
                    "success": False,
                    "file_path": file_path,
                    "output_path": target_path,
                    "saved": False,
                    "operation_count": len(parsed),
                    "completed_count": index,
                    "results": results,
                    "error": exc.to_dict(),
                }
            except Exception as exc:
                wrapped = OfficeOperationError("operation_failed", f"操作执行失败: {exc}", path=operation.path)
                results.append({"index": index, "success": False, "op": operation.op, "path": operation.path, "error": wrapped.to_dict()})
                return {
                    "success": False,
                    "file_path": file_path,
                    "output_path": target_path,
                    "saved": False,
                    "operation_count": len(parsed),
                    "completed_count": index,
                    "results": results,
                    "error": wrapped.to_dict(),
                }

        atomic_save(save, target_path)
        return {
            "success": True,
            "file_path": file_path,
            "output_path": target_path,
            "saved": True,
            "file_size": os.path.getsize(target_path),
            "operation_count": len(parsed),
            "completed_count": len(parsed),
            "results": results,
            "message": f"已原子执行 {len(parsed)} 个操作",
        }
    except OfficeOperationError as exc:
        return {"success": False, "file_path": file_path, "saved": False, "results": results, "error": exc.to_dict()}
    except ValueError as exc:
        error = OfficeOperationError("invalid_input", str(exc), path=file_path)
        return {"success": False, "file_path": file_path, "saved": False, "results": results, "error": error.to_dict()}
    except Exception as exc:
        error = OfficeOperationError("edit_failed", f"批量编辑失败: {exc}", path=file_path)
        return {"success": False, "file_path": file_path, "saved": False, "results": results, "error": error.to_dict()}
    finally:
        if workbook is not None:
            workbook.close()
