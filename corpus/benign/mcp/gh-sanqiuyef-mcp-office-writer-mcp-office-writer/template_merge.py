"""Placeholder-based template merge for DOCX, XLSX, and PPTX files."""
from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Any, Iterable

from .document_ops import OfficeOperationError
from .utils import atomic_save, ensure_output_dir, resolve_if_exists, validate_input_file, validate_output_path


_PLACEHOLDER = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _flatten_data(data: dict[str, Any]) -> dict[str, str]:
    values = {str(key): _stringify(value) for key, value in data.items()}

    def visit(value: Any, prefix: str) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                path = f"{prefix}.{key}" if prefix else str(key)
                values.setdefault(path, _stringify(child))
                visit(child, path)
        elif isinstance(value, list):
            for index, child in enumerate(value):
                path = f"{prefix}[{index}]"
                values.setdefault(path, _stringify(child))
                visit(child, path)

    for key, value in data.items():
        visit(value, str(key))
    return values


def _replace_placeholders(text: str, values: dict[str, str], used: set[str]) -> tuple[str, int]:
    replacements = 0

    def replace(match: re.Match[str]) -> str:
        nonlocal replacements
        key = match.group(1).strip()
        if key not in values:
            return match.group(0)
        used.add(key)
        replacements += 1
        return values[key]

    return _PLACEHOLDER.sub(replace, text), replacements


def _paragraph_text_runs(paragraph: Any) -> list[Any]:
    """返回段落中文档顺序的所有 w:r 元素（含 hyperlink 内部的 run）。

    ``paragraph.runs`` 只暴露段落直接子元素中的 run，hyperlink 内的 run 不在其中，
    但其文本会计入 ``paragraph.text``。占位符替换必须覆盖这些 run，
    否则 hyperlink 内的占位符会被计入替换次数却实际不生效。
    """
    from docx.oxml.ns import qn
    return list(paragraph._element.iter(qn("w:r")))


def _run_text(run_element: Any) -> str:
    from docx.oxml.ns import qn
    return "".join(t.text or "" for t in run_element.findall(qn("w:t")))


def _set_run_text(run_element: Any, text: str) -> None:
    """将文本写入 w:r 的第一个 w:t，合并并清空其余 w:t；没有 w:t 时新建。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    all_t = run_element.findall(qn("w:t"))
    if all_t:
        main_t = all_t[0]
        for extra in all_t[1:]:
            run_element.remove(extra)
    else:
        main_t = OxmlElement("w:t")
        run_element.append(main_t)
    main_t.text = text
    main_t.set(qn("xml:space"), "preserve")


def _replace_word_paragraph(paragraph: Any, values: dict[str, str], used: set[str]) -> int:
    text_runs = _paragraph_text_runs(paragraph)
    full_text = "".join(_run_text(r) for r in text_runs)
    changed, count = _replace_placeholders(full_text, values, used)
    if not count:
        return 0
    if text_runs:
        _set_run_text(text_runs[0], changed)
        for run_element in text_runs[1:]:
            for t in run_element.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
                t.text = ""
    else:
        paragraph.add_run(changed)
    return count


def _word_table_paragraphs(table: Any) -> Iterable[Any]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _word_table_paragraphs(nested)


def _word_paragraphs(document: Any) -> Iterable[Any]:
    seen: set[int] = set()

    def emit(paragraphs: Iterable[Any]) -> Iterable[Any]:
        for paragraph in paragraphs:
            marker = id(paragraph._element)
            if marker not in seen:
                seen.add(marker)
                yield paragraph

    yield from emit(document.paragraphs)
    for table in document.tables:
        yield from emit(_word_table_paragraphs(table))
    for section in document.sections:
        for container in (section.header, section.first_page_header, section.even_page_header,
                          section.footer, section.first_page_footer, section.even_page_footer):
            yield from emit(container.paragraphs)
            for table in container.tables:
                yield from emit(_word_table_paragraphs(table))


def _merge_docx(template_path: str, output_path: str, values: dict[str, str]) -> tuple[int, set[str], set[str]]:
    from docx import Document

    document = Document(template_path)
    used: set[str] = set()
    paragraphs = list(_word_paragraphs(document))
    count = sum(_replace_word_paragraph(paragraph, values, used) for paragraph in paragraphs)
    unresolved = {
        match.group(1).strip()
        for paragraph in paragraphs
        for match in _PLACEHOLDER.finditer(paragraph.text)
    }
    atomic_save(document.save, output_path)
    return count, used, unresolved


def _merge_xlsx(template_path: str, output_path: str, values: dict[str, str]) -> tuple[int, set[str], set[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(template_path, keep_vba=template_path.lower().endswith(".xlsm"))
    used: set[str] = set()
    count = 0
    try:
        for worksheet in workbook.worksheets:
            for cell in worksheet._cells.values():
                if not isinstance(cell.value, str):
                    continue
                changed, replacements = _replace_placeholders(cell.value, values, used)
                if replacements:
                    cell.value = changed
                    count += replacements
        unresolved = {
            match.group(1).strip()
            for worksheet in workbook.worksheets
            for cell in worksheet._cells.values()
            if isinstance(cell.value, str)
            for match in _PLACEHOLDER.finditer(cell.value)
        }
        atomic_save(workbook.save, output_path)
        return count, used, unresolved
    finally:
        workbook.close()


def _ppt_text_frames(presentation: Any) -> Iterable[Any]:
    seen: set[int] = set()
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                marker = id(shape.text_frame._txBody)
                if marker not in seen:
                    seen.add(marker)
                    yield shape.text_frame
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        marker = id(cell.text_frame._txBody)
                        if marker not in seen:
                            seen.add(marker)
                            yield cell.text_frame
        notes_frame = slide.notes_slide.notes_text_frame
        marker = id(notes_frame._txBody)
        if marker not in seen:
            seen.add(marker)
            yield notes_frame


def _merge_pptx(template_path: str, output_path: str, values: dict[str, str]) -> tuple[int, set[str], set[str]]:
    from pptx import Presentation

    presentation = Presentation(template_path)
    used: set[str] = set()
    frames = list(_ppt_text_frames(presentation))
    count = 0
    for frame in frames:
        for paragraph in frame.paragraphs:
            original = "".join(run.text for run in paragraph.runs)
            changed, replacements = _replace_placeholders(original, values, used)
            if not replacements:
                continue
            count += replacements
            if paragraph.runs:
                paragraph.runs[0].text = changed
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run().text = changed
    unresolved = {
        match.group(1).strip()
        for frame in frames
        for paragraph in frame.paragraphs
        for match in _PLACEHOLDER.finditer("".join(run.text for run in paragraph.runs))
    }
    atomic_save(presentation.save, output_path)
    return count, used, unresolved


def merge_office_template(
    template_path: str,
    output_path: str,
    data: dict[str, Any],
    if_exists: str = "overwrite",
) -> dict[str, Any]:
    """Fill {{key}} placeholders while preserving the surrounding Office template."""
    try:
        suffix = Path(template_path).suffix.lower()
        if suffix not in {".docx", ".xlsx", ".xlsm", ".pptx"}:
            raise OfficeOperationError(
                "unsupported_type",
                f"不支持的模板格式: {suffix or '(无扩展名)'}",
                path=template_path,
                valid_values=[".docx", ".xlsx", ".xlsm", ".pptx"],
            )
        validate_input_file(template_path, (suffix,))
        validate_output_path(output_path, (suffix,))
        ensure_output_dir(output_path)
        resolved_path = resolve_if_exists(output_path, if_exists)
        values = _flatten_data(data)
        if suffix == ".docx":
            count, used, unresolved = _merge_docx(template_path, resolved_path, values)
        elif suffix in {".xlsx", ".xlsm"}:
            count, used, unresolved = _merge_xlsx(template_path, resolved_path, values)
        else:
            count, used, unresolved = _merge_pptx(template_path, resolved_path, values)
        return {
            "success": True,
            "template_path": template_path,
            "file_path": resolved_path,
            "file_size": os.path.getsize(resolved_path),
            "replaced_count": count,
            "used_keys": sorted(used),
            "unresolved_placeholders": sorted(unresolved),
            "message": f"模板合并完成，替换 {count} 处，占位符未解析 {len(unresolved)} 个",
        }
    except OfficeOperationError as exc:
        return {"success": False, "template_path": template_path, "file_path": output_path, "error": exc.to_dict()}
    except (OSError, ValueError) as exc:
        error = OfficeOperationError("invalid_input", str(exc), path=template_path)
        return {"success": False, "template_path": template_path, "file_path": output_path, "error": error.to_dict()}
    except Exception as exc:
        error = OfficeOperationError("merge_failed", f"模板合并失败: {exc}", path=template_path)
        return {"success": False, "template_path": template_path, "file_path": output_path, "error": error.to_dict()}
