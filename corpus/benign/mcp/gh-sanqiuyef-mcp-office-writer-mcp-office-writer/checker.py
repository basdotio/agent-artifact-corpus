"""
mcp-office-writer 后处理验证模块
参考 document-skills 的 postcheck.py，对生成的 Office 文档进行质量检查
"""
from __future__ import annotations

import logging
import os
import re
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)


@dataclass
class CheckResult:
    """单项检查结果"""
    check_name: str
    passed: bool
    severity: str = "info"  # info / warn / error
    message: str = ""
    details: str = ""


@dataclass
class ValidationReport:
    """完整验证报告"""
    file_path: str
    checks: list[CheckResult] = field(default_factory=list)
    total: int = 0
    passed: int = 0
    warnings: int = 0
    errors: int = 0

    @property
    def summary(self) -> str:
        parts = [f"验证报告: {self.file_path}"]
        parts.append(f"  总计: {self.total} | 通过: {self.passed} | 警告: {self.warnings} | 错误: {self.errors}")
        for c in self.checks:
            icon = "✅" if c.passed else ("⚠️" if c.severity == "warn" else "❌")
            parts.append(f"  {icon} {c.check_name}: {c.message}")
        return "\n".join(parts)

    def to_dict(self) -> dict:
        return {
            "file_path": self.file_path,
            "total": self.total,
            "passed": self.passed,
            "warnings": self.warnings,
            "errors": self.errors,
            "checks": [
                {"name": c.check_name, "passed": c.passed, "severity": c.severity,
                 "message": c.message, "details": c.details}
                for c in self.checks
            ],
        }


def _add_check(report: ValidationReport, name: str, passed: bool,
               severity: str = "info", message: str = "", details: str = ""):
    """添加检查结果"""
    report.checks.append(CheckResult(name, passed, severity, message, details))
    report.total += 1
    if passed:
        report.passed += 1
    elif severity == "error":
        report.errors += 1
    else:
        report.warnings += 1


def _safe_line_spacing(pf) -> float | None:
    """安全提取段落行距数值，兼容枚举值"""
    ls = pf.line_spacing
    if ls is None:
        return None
    if isinstance(ls, (int, float)):
        return round(float(ls), 2)
    # 枚举值（如 WD_LINE_SPACING.ONE_POINT_FIVE），无法 round
    return None


def _has_page_break(run) -> bool:
    """检测 run 中是否包含分页符（安全 XML 检查）"""
    try:
        for child in run._element:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "br":
                if child.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page":
                    return True
        return False
    except Exception:
        return False


def _has_drawing(run) -> bool:
    """检测 run 中是否包含图片（安全 XML 检查）"""
    try:
        for child in run._element:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "drawing":
                return True
        return False
    except Exception:
        return False


def _get_heading_level(p) -> int | None:
    """安全提取段落标题层级"""
    style = p.style
    if not style or not style.name:
        return None
    m = re.search(r"(?:Heading|heading)\s*(\d+)", style.name)
    if m:
        return int(m.group(1))
    return None


def validate_word_document(file_path: str) -> ValidationReport:
    """对 Word 文档进行后处理质量检查"""
    report = ValidationReport(file_path=file_path)

    if not os.path.isfile(file_path):
        _add_check(report, "文件存在", False, "error", "文件不存在")
        return report

    try:
        from .utils import validate_input_file
        validate_input_file(file_path, (".docx",))
    except ValueError as e:
        _add_check(report, "输入文件安全检查", False, "error", str(e))
        return report

    try:
        from docx import Document

        doc = Document(file_path)

        # ── 检查 1: 文件基础信息 ──
        file_size = os.path.getsize(file_path)
        size_ok = file_size > 0
        _add_check(report, "文件大小检查", size_ok, "error" if not size_ok else "info",
                   f"{file_size} 字节" if size_ok else "空文件")

        para_count = len(doc.paragraphs)
        _add_check(report, "段落数量", True, "info", f"{para_count} 个段落")

        # ── 检查 2: 空白页检查 ──
        empty_para_count = sum(1 for p in doc.paragraphs if not p.text.strip())
        page_break_count = 0
        for p in doc.paragraphs:
            for run in p.runs:
                if _has_page_break(run):
                    page_break_count += 1

        if empty_para_count > 5 and page_break_count > 2:
            _add_check(report, "空白页检测", False, "warn",
                       f"检测到 {empty_para_count} 个空段落和 {page_break_count} 个分页符，可能存在空白页",
                       f"空段落: {empty_para_count}, 分页符: {page_break_count}")
        else:
            _add_check(report, "空白页检测", True, "info",
                       f"空段落: {empty_para_count}, 分页符: {page_break_count}")

        # ── 检查 3: 行间距一致性 ──
        line_spacings = set()
        for p in doc.paragraphs:
            pf = p.paragraph_format
            if pf.line_spacing is not None and p.text.strip():
                ls = _safe_line_spacing(pf)
                if ls is not None:
                    line_spacings.add(ls)
        if len(line_spacings) > 3:
            _add_check(report, "行间距一致性", False, "warn",
                       f"检测到 {len(line_spacings)} 种不同行间距: {sorted(line_spacings)}")
        else:
            _add_check(report, "行间距一致性", True, "info",
                       f"行间距: {sorted(line_spacings) if line_spacings else '默认'}")

        # ── 检查 4: 表格完整性 ──
        table_count = len(doc.tables)
        table_issues = 0
        for ti, table in enumerate(doc.tables):
            if len(table.rows) == 0:
                table_issues += 1
                continue
            has_content = False
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        has_content = True
                        break
            if not has_content:
                table_issues += 1
        if table_issues > 0:
            _add_check(report, "表格完整性", False, "warn",
                       f"共 {table_count} 个表格，{table_issues} 个可能存在空表问题")
        else:
            _add_check(report, "表格完整性", True, "info", f"共 {table_count} 个表格")

        # ── 检查 5: 标题层级连续（报告所有跳跃，不只是第一个）──
        heading_levels = []
        for p in doc.paragraphs:
            level = _get_heading_level(p)
            if level is not None:
                heading_levels.append((level, p.text.strip()[:40]))

        if heading_levels:
            jumps_found = False
            for i in range(1, len(heading_levels)):
                prev_lvl = heading_levels[i - 1][0]
                curr_lvl = heading_levels[i][0]
                if curr_lvl > prev_lvl + 1:
                    _add_check(report, "标题层级连续性", False, "warn",
                               f"标题从 H{prev_lvl} 跳转到 H{curr_lvl}（第 {i} 个标题）")
                    jumps_found = True
            if not jumps_found:
                levels_only = [level for level, _ in heading_levels]
                _add_check(report, "标题层级连续性", True, "info",
                           f"标题层级: H{min(levels_only)}-H{max(levels_only)}")
        else:
            _add_check(report, "标题层级连续性", True, "info", "无标题（可能为纯文档）")

        # ── 检查 6: 图片嵌入检测 ──
        image_count = 0
        for p in doc.paragraphs:
            for run in p.runs:
                if _has_drawing(run):
                    image_count += 1
        _add_check(report, "图片嵌入检查", True, "info", f"检测到 {image_count} 个图片")

        # ── 检查 7: section/分节设置 ──
        section_count = len(doc.sections)
        if section_count > 1:
            _add_check(report, "分节设置", True, "info", f"共 {section_count} 个节")
        else:
            _add_check(report, "分节设置", True, "info", "单节文档")

        # ── 检查 8: 页眉页脚 ──
        hf_found = False
        for section in doc.sections:
            if section.header.paragraphs and any(p.text.strip() for p in section.header.paragraphs):
                hf_found = True
                break
            if section.footer.paragraphs and any(p.text.strip() for p in section.footer.paragraphs):
                hf_found = True
                break
        if hf_found:
            _add_check(report, "页眉页脚", True, "info", "已检测到页眉页脚内容")
        else:
            _add_check(report, "页眉页脚", True, "info", "无页眉页脚（可能是封面/简历等）")

        # ── 检查 9: 遗留占位标记 ──
        marker_pattern = re.compile(r"\b(?:TODO|TBD|DRAFT)\b|待(?:完善|补充|确认)", re.IGNORECASE)
        text_parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            text_parts.extend(cell.text for row in table.rows for cell in row.cells)
        for section in doc.sections:
            text_parts.extend(p.text for p in section.header.paragraphs)
            text_parts.extend(p.text for p in section.footer.paragraphs)
        markers = sorted({match.group(0) for text in text_parts for match in marker_pattern.finditer(text)})
        if markers:
            _add_check(report, "遗留占位标记", False, "warn", f"发现: {', '.join(markers)}")
        else:
            _add_check(report, "遗留占位标记", True, "info", "未发现 TODO/TBD/DRAFT 或待完善标记")

    except Exception as e:
        _add_check(report, "文档解析", False, "error", f"无法解析文档: {e}")

    logger.info("文档验证完成: %s", report.summary.replace("\n", " | "))
    return report
